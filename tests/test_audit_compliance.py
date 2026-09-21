from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from klen_clone.audit_compliance import (audit_compliance_workspace_payload,
    generate_statutory_package, statutory_json_bytes, statutory_package_payload,
    statutory_workbook_bytes, transition_statutory_package)
from klen_clone.close_reporting import generate_report_package, transition_report_package
from klen_clone.operational import OperationalFiscalPeriod, initialize_operational_database, make_operational_engine
from klen_clone.period_close import (create_close_adjustment, create_period_close,
    decide_close_adjustment, rehearse_period_close, transition_period_close)


def approved_report():
    engine = make_operational_engine("sqlite:///:memory:")
    initialize_operational_database(engine)
    session = Session(engine)
    session.add(OperationalFiscalPeriod(period_key="2026-09", starts_on=date(2026, 9, 1),
        ends_on=date(2026, 9, 30), status="open", rehearsal_enabled=True,
        approval_reference="test", configured_by="controller")); session.commit()
    close = create_period_close(session, fiscal_period_key="2026-09", actor="maker")
    adjustment = create_close_adjustment(session, close, adjustment_type="accrual",
        adjustment_date=date(2026, 9, 30), debit_account="Utilities expense",
        credit_account="Accrued liabilities", amount="125.50", reversal_on=date(2026, 10, 1),
        evidence_reference="TEST-1", reason="September utility accrual", actor="maker")
    decide_close_adjustment(session, adjustment, action="approve", expected_revision=1,
        note="Approved evidence", actor="checker")
    transition_period_close(session, close, action="submit", expected_revision=1, actor="maker", note="Ready")
    transition_period_close(session, close, action="approve", expected_revision=2, actor="checker", note="Approved close")
    rehearse_period_close(session, close, actor="checker")
    report = generate_report_package(session, close, actor="report-maker")
    transition_report_package(session, report, action="submit", expected_revision=1, note="", actor="report-maker")
    transition_report_package(session, report, action="approve", expected_revision=2,
        note="Independent reporting approval", actor="report-checker")
    return session, report


def test_statutory_package_hash_chains_controls_and_exports_after_approval():
    session, report = approved_report()
    package = generate_statutory_package(session, report, actor="compliance-maker")
    payload = statutory_package_payload(package)
    assert payload["status"] == "draft"
    assert len(payload["evidence"]["audit_trail"]["chain_root"]) == 64
    assert payload["evidence"]["audit_trail"]["event_count"] > 0
    assert all(control["status"] != "fail" for control in payload["evidence"]["controls"])
    transition_statutory_package(session, package, action="submit", expected_revision=1,
        note="", actor="compliance-maker")
    with pytest.raises(PermissionError):
        transition_statutory_package(session, package, action="approve", expected_revision=2,
            note="Self approval is blocked", actor="compliance-maker")
    session.rollback()
    transition_statutory_package(session, package, action="approve", expected_revision=2,
        note="Independent compliance evidence approved", actor="compliance-checker")
    assert statutory_package_payload(package)["exports_enabled"] is True
    workbook = load_workbook(__import__("io").BytesIO(statutory_workbook_bytes(package)))
    assert {"Control Register", "Immutable Audit Trail", "VAT Evidence", "Manifest"}.issubset(workbook.sheetnames)
    assert b'"posting_enabled": false' in statutory_json_bytes(package)
    workspace = audit_compliance_workspace_payload(session)
    assert workspace["controls"]["approved"] == 1
    assert workspace["controls"]["permanent_postings"] == 0


def test_statutory_package_requires_approved_financial_report():
    session, report = approved_report()
    report.status = "submitted"; session.commit()
    with pytest.raises(ValueError, match="approved financial statements"):
        generate_statutory_package(session, report, actor="maker")
