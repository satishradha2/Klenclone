from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from klen_clone.cutover_rehearsal import (build_cutover_rehearsal,
    cutover_json_bytes, cutover_rehearsal_payload, cutover_workbook_bytes,
    generate_cutover_rehearsal, transition_cutover_rehearsal)
from klen_clone.operational import (OperationalAuditEvent, OperationalDraft,
    initialize_operational_database, make_operational_engine)


def session_with_test_data():
    engine = make_operational_engine("sqlite:///:memory:")
    initialize_operational_database(engine)
    session = Session(engine)
    session.add(OperationalDraft(draft_key="draft-1", draft_no="DR-S-1", document_type="sale",
        party_code="C1", party_name_snapshot="Test customer", location_code="MAIN",
        currency_code="AED", subtotal=10, discount_amount=0, tax_amount=0,
        total_amount=10, status="draft", posting_enabled=False, created_by="maker",
        state_changed_by="maker"))
    session.commit()
    return session


def test_rehearsal_inventories_without_purging_and_requires_checker():
    session = session_with_test_data()
    before_drafts = session.scalar(select(func.count(OperationalDraft.id)))
    plan = build_cutover_rehearsal(session)
    assert plan["execution_enabled"] is False
    assert plan["purge_performed"] is False and plan["import_performed"] is False
    assert any(item["table"] == "operational_drafts" and item["rows"] == 1 for item in plan["inventory"])
    assert any(item["status"] == "blocked" for item in plan["controls"])
    package = generate_cutover_rehearsal(session, actor="cutover-maker")
    assert generate_cutover_rehearsal(session, actor="cutover-maker").id == package.id
    assert session.scalar(select(func.count(OperationalDraft.id))) == before_drafts
    transition_cutover_rehearsal(session, package, action="submit", expected_revision=1,
        note="", actor="cutover-maker")
    with pytest.raises(PermissionError):
        transition_cutover_rehearsal(session, package, action="approve", expected_revision=2,
            note="Self approval blocked", actor="cutover-maker")
    session.rollback()
    transition_cutover_rehearsal(session, package, action="approve", expected_revision=2,
        note="Approved rehearsal plan only; no execution authority", actor="cutover-checker")
    payload = cutover_rehearsal_payload(package)
    assert payload["exports_enabled"] is True and payload["execution_enabled"] is False
    assert session.scalar(select(func.count(OperationalDraft.id))) == before_drafts
    assert session.scalar(select(func.count(OperationalAuditEvent.id))) == 3


def test_rehearsal_exports_include_gates_inventory_order_and_manifest():
    session = session_with_test_data(); package = generate_cutover_rehearsal(session, actor="maker")
    transition_cutover_rehearsal(session, package, action="submit", expected_revision=1, note="", actor="maker")
    transition_cutover_rehearsal(session, package, action="approve", expected_revision=2,
        note="Independent rehearsal approval", actor="checker")
    workbook = load_workbook(BytesIO(cutover_workbook_bytes(package)))
    assert {"Control Gates", "Table Inventory", "Purge Order", "Cutover Phases", "Manifest"}.issubset(workbook.sheetnames)
    manifest = cutover_json_bytes(package)
    assert b'"execution_enabled": false' in manifest
    assert b'"purge_performed": false' in manifest
    assert b'"source_mutated": false' in manifest
