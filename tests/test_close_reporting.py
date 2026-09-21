from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from klen_clone.close_reporting import (generate_report_package, pdf_bytes,
    report_package_payload, transition_report_package, workbook_bytes)
from klen_clone.operational import OperationalFiscalPeriod, initialize_operational_database, make_operational_engine
from klen_clone.period_close import (create_close_adjustment, create_period_close,
    decide_close_adjustment, rehearse_period_close, transition_period_close)


def approved_close():
    engine = make_operational_engine("sqlite:///:memory:"); initialize_operational_database(engine); session = Session(engine)
    session.add(OperationalFiscalPeriod(period_key="2026-09", starts_on=date(2026,9,1), ends_on=date(2026,9,30),
        status="open", rehearsal_enabled=True, approval_reference="test", configured_by="controller")); session.commit()
    close = create_period_close(session, fiscal_period_key="2026-09", actor="maker")
    item = create_close_adjustment(session, close, adjustment_type="accrual", adjustment_date=date(2026,9,30),
        debit_account="Utilities expense", credit_account="Accrued liabilities", amount="125.50",
        reversal_on=date(2026,10,1), evidence_reference="TEST-1", reason="September utility accrual", actor="maker")
    decide_close_adjustment(session, item, action="approve", expected_revision=1, note="Approved evidence", actor="checker")
    transition_period_close(session, close, action="submit", expected_revision=1, actor="maker", note="VAT warning test")
    transition_period_close(session, close, action="approve", expected_revision=2, actor="checker", note="Approved close")
    rehearse_period_close(session, close, actor="checker")
    return session, close


def test_close_report_is_balanced_controlled_and_exportable_after_approval():
    session, close = approved_close(); package = generate_report_package(session, close, actor="report-maker")
    report = report_package_payload(package)["report"]
    assert Decimal(str(report["trial_balance"]["debit"])) == Decimal("125.50")
    assert Decimal(str(report["trial_balance"]["credit"])) == Decimal("125.50")
    assert Decimal(str(report["profit_and_loss"]["profit"])) == Decimal("-125.50")
    assert Decimal(str(report["balance_sheet"]["difference"])) == Decimal("0.00")
    transition_report_package(session, package, action="submit", expected_revision=1, note="", actor="report-maker")
    with pytest.raises(PermissionError):
        transition_report_package(session, package, action="approve", expected_revision=2, note="Self approve", actor="report-maker")
    session.rollback()
    transition_report_package(session, package, action="approve", expected_revision=2,
        note="Independent reporting approval", actor="report-checker")
    payload = report_package_payload(package); assert payload["exports_enabled"] is True
    workbook = load_workbook(filename=__import__("io").BytesIO(workbook_bytes(package)))
    assert {"Trial Balance", "Profit and Loss", "Balance Sheet", "Cash Flow", "Retained Earnings"}.issubset(workbook.sheetnames)
    assert pdf_bytes(package).startswith(b"%PDF-1.4")


def test_reporting_requires_approved_close():
    session, close = approved_close(); close.status = "submitted"; session.commit()
    with pytest.raises(ValueError, match="approved frozen"):
        generate_report_package(session, close, actor="maker")
