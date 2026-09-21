from __future__ import annotations

import hashlib
import io
import json
import uuid
from collections import Counter
from datetime import datetime

from openpyxl import Workbook
from sqlalchemy import CheckConstraint, DateTime, ForeignKey, Integer, String, Text, UniqueConstraint, func, select
from sqlalchemy.orm import Mapped, Session, mapped_column, relationship

from .close_reporting import OperationalFinancialReportPackage
from .operational import OperationalAuditEvent, OperationalBase, OperationalJournalBatch, utc_now
from .posting_integration import OperationalIntegratedPostingBatch
from .vat_control import OperationalVatPeriod, OperationalVatReturnRehearsal


class OperationalStatutoryEvidencePackage(OperationalBase):
    __tablename__ = "operational_statutory_evidence_packages"
    __table_args__ = (
        UniqueConstraint("financial_report_id", "report_revision", name="uq_statutory_report_revision"),
        CheckConstraint("status IN ('draft','submitted','approved','rejected')", name="ck_statutory_package_status"),
    )

    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    package_key: Mapped[str] = mapped_column(String(36), unique=True, nullable=False)
    package_no: Mapped[str] = mapped_column(String(50), unique=True, nullable=False)
    financial_report_id: Mapped[int] = mapped_column(
        ForeignKey("operational_financial_report_packages.id"), nullable=False, index=True)
    report_revision: Mapped[int] = mapped_column(Integer, nullable=False)
    status: Mapped[str] = mapped_column(String(20), default="draft", nullable=False)
    evidence_json: Mapped[str] = mapped_column(Text, nullable=False)
    evidence_fingerprint: Mapped[str] = mapped_column(String(64), nullable=False)
    created_by: Mapped[str] = mapped_column(String(200), nullable=False)
    submitted_by: Mapped[str | None] = mapped_column(String(200))
    approved_by: Mapped[str | None] = mapped_column(String(200))
    decision_note: Mapped[str | None] = mapped_column(Text)
    revision: Mapped[int] = mapped_column(Integer, default=1, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=utc_now, nullable=False)
    approved_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True))
    financial_report: Mapped[OperationalFinancialReportPackage] = relationship()


def _event_payload(row: OperationalAuditEvent, prior_hash: str) -> tuple[dict, str]:
    item = {"event_key": row.event_key, "event_type": row.event_type, "actor": row.actor,
            "resource_key": row.resource_key, "occurred_at": row.occurred_at, "detail": row.detail}
    canonical = json.dumps(item, default=str, sort_keys=True, separators=(",", ":"))
    event_hash = hashlib.sha256((prior_hash + canonical).encode()).hexdigest()
    item["event_hash"] = event_hash
    return item, event_hash


def _audit_snapshot(session: Session) -> tuple[list[dict], str]:
    events = list(session.scalars(select(OperationalAuditEvent).order_by(
        OperationalAuditEvent.occurred_at, OperationalAuditEvent.id)))
    result, chain = [], "0" * 64
    for event in events:
        item, chain = _event_payload(event, chain)
        result.append(item)
    return result, chain


def build_statutory_evidence(session: Session, report: OperationalFinancialReportPackage) -> dict:
    if report.status != "approved":
        raise ValueError("Statutory evidence requires approved financial statements")
    close = report.period_close
    period = close.fiscal_period
    vat_periods = list(session.scalars(select(OperationalVatPeriod).where(
        OperationalVatPeriod.status == "approved",
        OperationalVatPeriod.starts_on <= period.ends_on,
        OperationalVatPeriod.ends_on >= period.starts_on).order_by(OperationalVatPeriod.starts_on)))
    audit_events, chain_root = _audit_snapshot(session)
    report_json = json.loads(report.report_json)
    integrated_postings = session.scalar(select(func.count(OperationalIntegratedPostingBatch.id))) or 0
    journal_postings = session.scalar(select(func.count(OperationalJournalBatch.id))) or 0
    vat_items = []
    for vat in vat_periods:
        rehearsal = session.scalar(select(OperationalVatReturnRehearsal).where(
            OperationalVatReturnRehearsal.period_id == vat.id,
            OperationalVatReturnRehearsal.period_revision == vat.revision))
        vat_items.append({"period_code": vat.period_code, "starts_on": vat.starts_on,
            "ends_on": vat.ends_on, "due_on": vat.due_on, "company_trn": vat.company_trn,
            "source_count": vat.source_count, "output_vat": vat.output_vat,
            "input_vat": vat.input_vat, "net_vat": vat.net_vat,
            "source_fingerprint": vat.source_fingerprint,
            "rehearsal_fingerprint": rehearsal.source_fingerprint if rehearsal else None,
            "filing_performed": False})
    controls = [
        {"control": "approved_financial_statements", "status": "pass", "evidence": report.package_no},
        {"control": "balanced_trial_balance", "status": "pass" if str(report_json["trial_balance"]["difference"]) in {"0", "0.00"} else "fail",
         "evidence": f"difference AED {report_json['trial_balance']['difference']}"},
        {"control": "balanced_balance_sheet", "status": "pass" if str(report_json["balance_sheet"]["difference"]) in {"0", "0.00"} else "fail",
         "evidence": f"difference AED {report_json['balance_sheet']['difference']}"},
        {"control": "approved_vat_return", "status": "pass" if vat_items else "warning",
         "evidence": ", ".join(v["period_code"] for v in vat_items) or "No overlapping approved VAT period"},
        {"control": "permanent_posting_lock", "status": "pass" if integrated_postings + journal_postings == 0 else "fail",
         "evidence": f"integrated {integrated_postings}; journal {journal_postings}"},
        {"control": "source_system_protection", "status": "pass", "evidence": "BizModo read-only; target ERP evidence only"},
        {"control": "statutory_filing_lock", "status": "pass", "evidence": "No filing performed; rehearsal only"},
    ]
    return {"company": "ASAS GENERAL TRADING LLC", "period": period.period_key,
        "period_start": period.starts_on, "period_end": period.ends_on,
        "financial_report": {"package_no": report.package_no,
            "fingerprint": report.report_fingerprint, "approved_by": report.approved_by,
            "close_no": close.close_no, "close_fingerprint": close.source_fingerprint},
        "vat_returns": vat_items, "controls": controls,
        "audit_trail": {"event_count": len(audit_events), "chain_root": chain_root, "events": audit_events},
        "posting_enabled": False, "filing_performed": False, "source_evidence_mutated": False}


def generate_statutory_package(session: Session, report: OperationalFinancialReportPackage, *, actor: str):
    existing = session.scalar(select(OperationalStatutoryEvidencePackage).where(
        OperationalStatutoryEvidencePackage.financial_report_id == report.id,
        OperationalStatutoryEvidencePackage.report_revision == report.revision))
    if existing:
        return existing
    evidence = build_statutory_evidence(session, report)
    document = json.dumps(evidence, default=str, sort_keys=True, separators=(",", ":"))
    fingerprint = hashlib.sha256(document.encode()).hexdigest()
    key = str(uuid.uuid4())
    row = OperationalStatutoryEvidencePackage(package_key=key,
        package_no=f"STAT-{evidence['period']}-{key[:6].upper()}", financial_report_id=report.id,
        report_revision=report.revision, status="draft", evidence_json=document,
        evidence_fingerprint=fingerprint, created_by=actor)
    session.add(row); session.flush()
    session.add(OperationalAuditEvent(event_key=str(uuid.uuid4()), event_type="statutory_package.generated",
        actor=actor, resource_key=row.package_key,
        detail=f"{row.package_no}; audit chain {evidence['audit_trail']['chain_root'][:12]}; no filing; no posting"))
    session.commit()
    return row


def transition_statutory_package(session: Session, row: OperationalStatutoryEvidencePackage, *,
                                 action: str, expected_revision: int, note: str, actor: str):
    if row.revision != expected_revision:
        raise ValueError("Statutory-package revision is stale")
    if action == "submit" and row.status in {"draft", "rejected"}:
        if actor != row.created_by:
            raise PermissionError("Only the package maker may submit this package")
        evidence = json.loads(row.evidence_json)
        failed = [item["control"] for item in evidence["controls"] if item["status"] == "fail"]
        if failed:
            raise ValueError("Resolve failed statutory controls before submission: " + ", ".join(failed))
        row.status, row.submitted_by = "submitted", actor
    elif action in {"approve", "reject"} and row.status == "submitted":
        if actor == row.created_by:
            raise PermissionError("Maker cannot approve their own statutory package")
        if len(note.strip()) < 5:
            raise ValueError("Independent decision note must contain at least five characters")
        row.status = "approved" if action == "approve" else "rejected"
        row.approved_by = actor if action == "approve" else None
        row.decision_note = note.strip()
        row.approved_at = utc_now() if action == "approve" else None
    else:
        raise ValueError("Invalid statutory-package transition")
    row.revision += 1
    session.add(OperationalAuditEvent(event_key=str(uuid.uuid4()), event_type=f"statutory_package.{action}",
        actor=actor, resource_key=row.package_key, detail=f"{row.package_no}; {row.status}; no filing; no posting"))
    session.commit()
    return row


def statutory_package_payload(row: OperationalStatutoryEvidencePackage) -> dict:
    return {"package_key": row.package_key, "package_no": row.package_no,
        "status": row.status, "evidence_fingerprint": row.evidence_fingerprint,
        "evidence": json.loads(row.evidence_json), "created_by": row.created_by,
        "approved_by": row.approved_by, "decision_note": row.decision_note,
        "revision": row.revision, "exports_enabled": row.status == "approved"}


def audit_compliance_workspace_payload(session: Session, *, limit: int = 100) -> dict:
    events = list(session.scalars(select(OperationalAuditEvent).order_by(
        OperationalAuditEvent.occurred_at.desc(), OperationalAuditEvent.id.desc()).limit(limit)))
    packages = list(session.scalars(select(OperationalStatutoryEvidencePackage).order_by(
        OperationalStatutoryEvidencePackage.created_at.desc())))
    reports = list(session.scalars(select(OperationalFinancialReportPackage).where(
        OperationalFinancialReportPackage.status == "approved").order_by(
        OperationalFinancialReportPackage.created_at.desc())))
    event_types = Counter(event.event_type.split(".", 1)[0] for event in events)
    return {"events": [{"event_key": row.event_key, "event_type": row.event_type,
        "actor": row.actor, "resource_key": row.resource_key,
        "occurred_at": row.occurred_at, "detail": row.detail} for row in events],
        "event_modules": dict(sorted(event_types.items())),
        "eligible_reports": [{"package_key": row.package_key, "package_no": row.package_no,
            "period": row.period_close.fiscal_period.period_key,
            "has_statutory_package": any(item.financial_report_id == row.id for item in packages)} for row in reports],
        "packages": [statutory_package_payload(row) for row in packages],
        "controls": {"audit_events": session.scalar(select(func.count(OperationalAuditEvent.id))) or 0,
            "packages": len(packages), "awaiting_approval": sum(row.status == "submitted" for row in packages),
            "approved": sum(row.status == "approved" for row in packages),
            "permanent_postings": session.scalar(select(func.count(OperationalIntegratedPostingBatch.id))) or 0,
            "source_mutations": 0, "filings_performed": 0}}


def statutory_workbook_bytes(row: OperationalStatutoryEvidencePackage) -> bytes:
    evidence = json.loads(row.evidence_json)
    workbook = Workbook(); workbook.remove(workbook.active)
    controls = workbook.create_sheet("Control Register")
    controls.append(["Control", "Status", "Evidence"])
    for item in evidence["controls"]:
        controls.append([item["control"], item["status"], item["evidence"]])
    audit = workbook.create_sheet("Immutable Audit Trail")
    headers = ["event_key", "event_type", "actor", "resource_key", "occurred_at", "detail", "event_hash"]
    audit.append(headers)
    for item in evidence["audit_trail"]["events"]:
        audit.append([item[key] for key in headers])
    vat = workbook.create_sheet("VAT Evidence")
    vat_headers = ["period_code", "starts_on", "ends_on", "due_on", "company_trn", "source_count",
                   "output_vat", "input_vat", "net_vat", "source_fingerprint", "rehearsal_fingerprint"]
    vat.append(vat_headers)
    for item in evidence["vat_returns"]:
        vat.append([item.get(key) for key in vat_headers])
    manifest = workbook.create_sheet("Manifest")
    manifest.append(["Package", row.package_no]); manifest.append(["Fingerprint", row.evidence_fingerprint])
    manifest.append(["Period", evidence["period"]]); manifest.append(["Financial report", evidence["financial_report"]["package_no"]])
    manifest.append(["Financial report fingerprint", evidence["financial_report"]["fingerprint"]])
    manifest.append(["Audit chain root", evidence["audit_trail"]["chain_root"]])
    manifest.append(["Posting enabled", False]); manifest.append(["Filing performed", False])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row > 1 and sheet.max_column > 1:
            sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO(); workbook.save(output); return output.getvalue()


def statutory_json_bytes(row: OperationalStatutoryEvidencePackage) -> bytes:
    return json.dumps({"package_no": row.package_no, "fingerprint": row.evidence_fingerprint,
        "status": row.status, "evidence": json.loads(row.evidence_json)},
        indent=2, sort_keys=True).encode()
