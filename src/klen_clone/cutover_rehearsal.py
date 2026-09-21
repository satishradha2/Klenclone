from __future__ import annotations

import hashlib
import io
import json
import uuid
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from sqlalchemy import CheckConstraint, DateTime, Integer, String, Text, UniqueConstraint, func, inspect, select, text
from sqlalchemy.orm import Mapped, Session, mapped_column

from .operational import OperationalAuditEvent, OperationalBase, OperationalJournalBatch, utc_now
from .posting_integration import OperationalIntegratedPostingBatch
from .source_verification import OperationalSourceVerificationBatch


class OperationalCutoverRehearsalPackage(OperationalBase):
    __tablename__ = "operational_cutover_rehearsal_packages"
    __table_args__ = (
        UniqueConstraint("inventory_fingerprint", name="uq_cutover_rehearsal_inventory"),
        CheckConstraint("status IN ('draft','submitted','approved','rejected')", name="ck_cutover_rehearsal_status"),
        CheckConstraint("execution_enabled = false", name="ck_cutover_rehearsal_no_execution"),
    )

    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    package_key: Mapped[str] = mapped_column(String(36), unique=True, nullable=False)
    package_no: Mapped[str] = mapped_column(String(60), unique=True, nullable=False)
    status: Mapped[str] = mapped_column(String(20), default="draft", nullable=False)
    plan_json: Mapped[str] = mapped_column(Text, nullable=False)
    inventory_fingerprint: Mapped[str] = mapped_column(String(64), nullable=False)
    created_by: Mapped[str] = mapped_column(String(200), nullable=False)
    submitted_by: Mapped[str | None] = mapped_column(String(200))
    approved_by: Mapped[str | None] = mapped_column(String(200))
    decision_note: Mapped[str | None] = mapped_column(Text)
    revision: Mapped[int] = mapped_column(Integer, default=1, nullable=False)
    execution_enabled: Mapped[bool] = mapped_column(default=False, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=utc_now, nullable=False)
    approved_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True))


PRESERVE_UNTIL_ARCHIVED = {
    "operational_schema_migrations",
    "operational_cutover_rehearsal_packages",
}


def _table_inventory(session: Session) -> tuple[list[dict], list[str]]:
    inspector = inspect(session.bind)
    tables = sorted(name for name in inspector.get_table_names() if name.startswith("operational_"))
    counts = {}
    for name in tables:
        quoted = session.bind.dialect.identifier_preparer.quote(name)
        counts[name] = int(session.execute(text(f"SELECT COUNT(*) FROM {quoted}")).scalar_one())
    parents = defaultdict(set)
    children = defaultdict(set)
    for table in tables:
        for fk in inspector.get_foreign_keys(table):
            parent = fk.get("referred_table")
            if parent in tables and parent != table:
                parents[table].add(parent)
                children[parent].add(table)
    remaining, order = set(tables), []
    while remaining:
        leaves = sorted(name for name in remaining if not (children[name] & remaining))
        if not leaves:
            leaves = [sorted(remaining)[0]]
        order.extend(leaves)
        remaining.difference_update(leaves)
    inventory = [{"table": name, "rows": counts[name],
        "classification": "retain_until_external_archive" if name in PRESERVE_UNTIL_ARCHIVED else "purge_test_data",
        "depends_on": sorted(parents[name])} for name in tables]
    return inventory, order


def build_cutover_rehearsal(session: Session) -> dict:
    inventory, dependency_order = _table_inventory(session)
    purge_tables = [item for item in inventory if item["classification"] == "purge_test_data"]
    source_batches = list(session.scalars(select(OperationalSourceVerificationBatch).where(
        OperationalSourceVerificationBatch.status == "verified").order_by(
        OperationalSourceVerificationBatch.verified_at.desc())))
    integrated = session.scalar(select(func.count(OperationalIntegratedPostingBatch.id))) or 0
    journals = session.scalar(select(func.count(OperationalJournalBatch.id))) or 0
    controls = [
        {"control": "posting_lock", "status": "pass" if integrated + journals == 0 else "fail",
         "evidence": f"integrated {integrated}; permanent journals {journals}"},
        {"control": "source_read_only", "status": "pass",
         "evidence": "No cutover endpoint writes to BizModo"},
        {"control": "operational_inventory", "status": "pass",
         "evidence": f"{len(purge_tables)} purge tables; {sum(item['rows'] for item in purge_tables)} rows inventoried"},
        {"control": "final_frozen_bizmodo_capture", "status": "blocked",
         "evidence": "Current verified capture is test evidence; final transaction-free export not supplied"},
        {"control": "uploaded_file_archive", "status": "blocked",
         "evidence": "Complete BizModo attachment archive not yet supplied"},
        {"control": "target_backup_restore_test", "status": "blocked",
         "evidence": "Pre-purge target backup and restore checksum test must be recorded"},
        {"control": "final_row_count_manifest", "status": "blocked",
         "evidence": "Fresh final export counts and checksums must be approved"},
        {"control": "business_cutover_authorization", "status": "blocked",
         "evidence": "A separate named execution approval is required after ERP completion"},
    ]
    phases = [
        {"sequence": 1, "phase": "freeze_source", "action": "Declare a transaction-free BizModo window and capture final read-only exports plus checksums"},
        {"sequence": 2, "phase": "protect_target", "action": "Stop target writes; create and restore-test an encrypted target database backup"},
        {"sequence": 3, "phase": "archive_controls", "action": "Export approved audit, statutory and cutover packages outside the database"},
        {"sequence": 4, "phase": "purge_test_data", "action": "Delete target operational test rows child-first using the approved table plan"},
        {"sequence": 5, "phase": "fresh_import", "action": "Import final BizModo masters, opening balances, stock, transactions, attachments and HRM; payroll excluded"},
        {"sequence": 6, "phase": "reconcile", "action": "Reconcile row counts, checksums, stock, AR/AP, VAT and trial balance to the final manifest"},
        {"sequence": 7, "phase": "business_acceptance", "action": "Obtain independent business and finance approval; keep posting disabled"},
        {"sequence": 8, "phase": "activation", "action": "Enable production and posting only under a separate explicit authorization"},
    ]
    return {"mode": "non_destructive_cutover_reset_rehearsal",
        "execution_enabled": False, "purge_performed": False, "import_performed": False,
        "source_mutated": False, "posting_enabled": False,
        "inventory": inventory, "purge_order": [name for name in dependency_order if name not in PRESERVE_UNTIL_ARCHIVED],
        "archive_then_purge_last": sorted(PRESERVE_UNTIL_ARCHIVED),
        "summary": {"operational_tables": len(inventory), "purge_tables": len(purge_tables),
            "test_rows": sum(item["rows"] for item in purge_tables),
            "retained_control_rows": sum(item["rows"] for item in inventory if item["classification"] != "purge_test_data"),
            "verified_test_source_batches": len(source_batches)},
        "source_evidence": [{"batch_key": row.batch_key, "snapshot": row.source_snapshot_name,
            "records": row.verified_records, "checksum": row.evidence_checksum,
            "classification": "test_only_not_final_cutover"} for row in source_batches],
        "controls": controls, "phases": phases,
        "rollback": {"trigger": "Any failed checksum, balance, stock, VAT, attachment or business acceptance control",
            "action": "Keep target isolated, restore the verified pre-purge target backup, retain the rejected import evidence, and restart from a new final source capture",
            "bizmodo_action": "None; BizModo remains read-only"}}


def generate_cutover_rehearsal(session: Session, *, actor: str):
    active = session.scalar(select(OperationalCutoverRehearsalPackage).where(
        OperationalCutoverRehearsalPackage.status.in_({"draft", "submitted", "approved"})).order_by(
        OperationalCutoverRehearsalPackage.created_at.desc()))
    if active:
        return active
    plan = build_cutover_rehearsal(session)
    document = json.dumps(plan, default=str, sort_keys=True, separators=(",", ":"))
    fingerprint = hashlib.sha256(document.encode()).hexdigest()
    existing = session.scalar(select(OperationalCutoverRehearsalPackage).where(
        OperationalCutoverRehearsalPackage.inventory_fingerprint == fingerprint))
    if existing:
        return existing
    key = str(uuid.uuid4())
    row = OperationalCutoverRehearsalPackage(package_key=key,
        package_no=f"CUTOVER-REHEARSAL-{key[:8].upper()}", status="draft",
        plan_json=document, inventory_fingerprint=fingerprint, created_by=actor,
        execution_enabled=False)
    session.add(row); session.flush()
    session.add(OperationalAuditEvent(event_key=str(uuid.uuid4()), event_type="cutover.rehearsal_generated",
        actor=actor, resource_key=row.package_key,
        detail=f"{row.package_no}; {plan['summary']['test_rows']} test rows inventoried; no purge; no import"))
    session.commit()
    return row


def transition_cutover_rehearsal(session: Session, row: OperationalCutoverRehearsalPackage, *,
                                 action: str, expected_revision: int, note: str, actor: str):
    if row.revision != expected_revision:
        raise ValueError("Cutover-rehearsal revision is stale")
    if action == "submit" and row.status in {"draft", "rejected"}:
        if actor != row.created_by:
            raise PermissionError("Only the rehearsal maker may submit this package")
        failed = [item["control"] for item in json.loads(row.plan_json)["controls"] if item["status"] == "fail"]
        if failed:
            raise ValueError("Resolve failed safety controls before submission: " + ", ".join(failed))
        row.status, row.submitted_by = "submitted", actor
    elif action in {"approve", "reject"} and row.status == "submitted":
        if actor == row.created_by:
            raise PermissionError("Maker cannot approve their own cutover rehearsal")
        if len(note.strip()) < 5:
            raise ValueError("Independent decision note must contain at least five characters")
        row.status = "approved" if action == "approve" else "rejected"
        row.approved_by = actor if action == "approve" else None
        row.decision_note = note.strip()
        row.approved_at = utc_now() if action == "approve" else None
    else:
        raise ValueError("Invalid cutover-rehearsal transition")
    row.revision += 1
    session.add(OperationalAuditEvent(event_key=str(uuid.uuid4()), event_type=f"cutover.rehearsal_{action}",
        actor=actor, resource_key=row.package_key,
        detail=f"{row.package_no}; {row.status}; execution disabled; no purge; no import"))
    session.commit()
    return row


def cutover_rehearsal_payload(row: OperationalCutoverRehearsalPackage) -> dict:
    return {"package_key": row.package_key, "package_no": row.package_no,
        "status": row.status, "inventory_fingerprint": row.inventory_fingerprint,
        "plan": json.loads(row.plan_json), "created_by": row.created_by,
        "approved_by": row.approved_by, "decision_note": row.decision_note,
        "revision": row.revision, "execution_enabled": False,
        "exports_enabled": row.status == "approved"}


def cutover_rehearsal_workspace_payload(session: Session) -> dict:
    packages = list(session.scalars(select(OperationalCutoverRehearsalPackage).order_by(
        OperationalCutoverRehearsalPackage.created_at.desc())))
    live_plan = build_cutover_rehearsal(session)
    return {"live_plan": live_plan, "packages": [cutover_rehearsal_payload(row) for row in packages],
        "controls": {"packages": len(packages), "awaiting_approval": sum(row.status == "submitted" for row in packages),
            "approved": sum(row.status == "approved" for row in packages), "execution_enabled": False,
            "purges_performed": 0, "imports_performed": 0}}


def cutover_workbook_bytes(row: OperationalCutoverRehearsalPackage) -> bytes:
    plan = json.loads(row.plan_json); workbook = Workbook(); workbook.remove(workbook.active)
    sheets = {
        "Control Gates": (plan["controls"], ["control", "status", "evidence"]),
        "Table Inventory": (plan["inventory"], ["table", "rows", "classification", "depends_on"]),
        "Purge Order": ([{"sequence": index, "table": name} for index, name in enumerate(plan["purge_order"], 1)], ["sequence", "table"]),
        "Cutover Phases": (plan["phases"], ["sequence", "phase", "action"]),
    }
    for title, (rows, headers) in sheets.items():
        sheet = workbook.create_sheet(title); sheet.append(headers)
        for item in rows:
            sheet.append([json.dumps(item.get(key)) if isinstance(item.get(key), (list, dict)) else item.get(key) for key in headers])
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    manifest = workbook.create_sheet("Manifest")
    manifest.append(["Package", row.package_no]); manifest.append(["Fingerprint", row.inventory_fingerprint])
    manifest.append(["Execution enabled", False]); manifest.append(["Purge performed", False])
    manifest.append(["Import performed", False]); manifest.append(["Source mutated", False])
    output = io.BytesIO(); workbook.save(output); return output.getvalue()


def cutover_json_bytes(row: OperationalCutoverRehearsalPackage) -> bytes:
    return json.dumps({"package_no": row.package_no, "fingerprint": row.inventory_fingerprint,
        "status": row.status, "plan": json.loads(row.plan_json)}, indent=2, sort_keys=True).encode()
