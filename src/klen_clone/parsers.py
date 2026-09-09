from __future__ import annotations

import csv
import hashlib
import json
import mimetypes
import re
import zipfile
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree


ENTITY_RULES: tuple[tuple[str, str], ...] = (
    ("product_sales_detail_boundary_", "sales_line_boundary"),
    ("product_sales_detail_", "sales_line"),
    ("product_purchase_detail_", "purchase_line"),
    ("cash_flow_boundary_", "cash_flow_boundary"),
    ("cash_flow_", "cash_flow"),
    ("sales_return_details", "sales_return_detail"),
    ("purchase_return_details", "purchase_return_detail"),
    ("stock_transfer_details_", "stock_transfer_detail"),
    ("role_permissions_", "role_permission_matrix"),
    ("business_settings_redacted", "business_setting"),
    ("configuration_inventories", "configuration_inventory"),
    ("sales_payments_", "sales_payment"),
    ("purchase_payments_", "purchase_payment"),
    ("sales_returns_", "sales_return"),
    ("purchase_returns_", "purchase_return"),
    ("pos_sales_", "pos_sale"),
    ("sales_drafts_", "sales_draft"),
    ("sales_quotations_", "sales_quotation"),
    ("sales_", "sale"),
    ("purchases_", "purchase"),
    ("stock_transfers_", "stock_transfer"),
    ("stock_snapshot_", "stock_balance"),
    ("item_traceability_", "item_trace"),
    ("tax_input_", "tax_input"),
    ("tax_output_", "tax_output"),
    ("profit_by_product_", "profit_by_product"),
    ("trial_balance_", "trial_balance"),
)

DOCUMENT_KEYS = (
    "Invoice No.", "Invoice No", "Reference No", "Reference No.",
    "Purchase No", "Purchase Reference No", "Purchase Reference No.", "Parent Sale", "Parent Purchase",
)
ID_KEYS = ("ID", "Id", "id", "Transaction ID", "Transaction Id")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def infer_entity(path: Path) -> str:
    name = path.stem.lower()
    for prefix, entity in ENTITY_RULES:
        if name.startswith(prefix):
            return entity
    aliases = {
        "customers_all": "customer", "suppliers_all": "supplier", "products_all": "product",
        "users": "user", "units": "unit", "brands": "brand", "product_categories": "category",
        "expense_categories": "expense_category", "payment_accounts": "payment_account",
        "zones": "zone", "industries": "industry", "crm_contact_logins": "contact_login",
        "hrm_attendance_2026": "attendance", "hrm_shifts": "shift",
        "hrm_sales_targets": "sales_target", "shipments_2026": "shipment",
    }
    return aliases.get(name, name)


def media_type(path: Path) -> str:
    return mimetypes.guess_type(path.name)[0] or "application/octet-stream"


def read_records(path: Path) -> Iterable[dict[str, Any] | list[Any] | str]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            yield from csv.DictReader(stream)
        return
    if suffix == ".json":
        value = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(value, list):
            yield from value
        else:
            yield value
        return
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except (TypeError, ValueError):
            yield from read_xlsx_xml(path)
        else:
            try:
                sheet = workbook.active
                rows = sheet.iter_rows(values_only=True)
                headers = [str(value or "").strip() for value in next(rows, ())]
                for values in rows:
                    yield {headers[i] or f"column_{i + 1}": value for i, value in enumerate(values)}
            finally:
                workbook.close()
        return
    yield path.read_text(encoding="utf-8", errors="replace")


def read_xlsx_xml(path: Path) -> Iterable[dict[str, Any]]:
    """Read values from a malformed-styles XLSX without modifying the workbook."""
    namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as archive:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = ["".join(node.text or "" for node in item.findall(".//x:t", namespace)) for item in root.findall("x:si", namespace)]
        sheets = sorted(name for name in archive.namelist() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name))
        if not sheets:
            return
        root = ElementTree.fromstring(archive.read(sheets[0]))
        rows: list[list[Any]] = []
        for row in root.findall(".//x:sheetData/x:row", namespace):
            values: dict[int, Any] = {}
            for cell in row.findall("x:c", namespace):
                reference = cell.get("r", "A1")
                letters = re.match(r"[A-Z]+", reference).group(0)
                column = 0
                for letter in letters:
                    column = column * 26 + ord(letter) - 64
                kind = cell.get("t")
                value_node = cell.find("x:v", namespace)
                inline = cell.find("x:is", namespace)
                raw = value_node.text if value_node is not None else None
                if kind == "s" and raw is not None:
                    value: Any = shared[int(raw)]
                elif kind == "inlineStr" and inline is not None:
                    value = "".join(node.text or "" for node in inline.findall(".//x:t", namespace))
                elif raw is None:
                    value = None
                else:
                    try:
                        value = float(raw) if "." in raw else int(raw)
                    except ValueError:
                        value = raw
                values[column - 1] = value
            width = max(values, default=-1) + 1
            rows.append([values.get(index) for index in range(width)])
        if not rows:
            return
        headers = [str(value or "").strip() for value in rows[0]]
        for values in rows[1:]:
            yield {headers[i] if i < len(headers) and headers[i] else f"column_{i + 1}": value for i, value in enumerate(values)}


def presentation_row(record: Any) -> bool:
    if not isinstance(record, dict):
        return False
    values = [str(value or "").strip() for value in record.values() if not isinstance(value, (dict, list))]
    joined = " ".join(values).lower()
    first = next((value for value in values if value), "").lower()
    return first in {"total:", "total", "account"} or any(marker in joined for marker in ("delete selected", "select all", "bulk edit"))


def source_keys(record: Any) -> tuple[str | None, str | None]:
    if not isinstance(record, dict):
        return None, None
    source_id = next((str(record[key]).strip() for key in ID_KEYS if record.get(key)), None)
    document = next((str(record[key]).strip() for key in DOCUMENT_KEYS if record.get(key)), None)
    if not document:
        text = str(record.get("text") or record.get("snapshot") or "")
        match = re.search(r"(?:Invoice No\.:|Reference No: #?)([^\n)]+)", text)
        document = match.group(1).strip() if match else None
    return source_id, document
